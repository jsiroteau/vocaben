"""Outil de travail pour vocab.xlsx (base VocabEN).

Usage en ligne de commande (depuis la racine du dépôt) :
    python3 tools/vocab.py show <entrée> [<entrée>...]   affiche une ou plusieurs fiches
    python3 tools/vocab.py grep <regex>                    lignes dont Mot/Traduction/Synonymes correspondent
    python3 tools/vocab.py check                           contrôles automatiques

Usage en Python :
    import sys; sys.path.insert(0, 'tools'); from vocab import *
    B = Base()                       # charge vocab.xlsx
    r = B.find('to grieve', 'to grieve')          # une ligne (liste modifiable)
    r[T] = 'nouvelle traduction'; B.setex(r, [(ex, mot, trad), ...])
    B.insert_after(r, B.row('to grieve', 'to grieve sb', 'verb', 'peiner qn', ex=[...], zipf_of='grieve'))
    B.save()                         # contrôle puis enregistre (refuse si erreur)
"""
import re
import sys

import openpyxl
from openpyxl.styles import Font

FILE = 'vocab.xlsx'
# Index des colonnes
E, M, NAT, T, SYN, ANT, FA = range(7)
FREQ, ZIPF, TYPE = 22, 23, 24
NCOL = 25


def zipf(word):
    try:
        from wordfreq import zipf_frequency
    except ImportError:
        import subprocess
        subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', 'wordfreq', '--break-system-packages'])
        from wordfreq import zipf_frequency
    return round(zipf_frequency(word, 'en'), 2)


def freq_class(z):
    return 1 if z >= 4.87 else 2 if z >= 4.16 else 3 if z >= 3.5 else 4


def getex(r):
    return [(r[7 + 3 * k], r[8 + 3 * k], r[9 + 3 * k]) for k in range(5) if r[7 + 3 * k]]


def setex(r, ex):
    assert len(ex) <= 5, 'cinq exemples au plus'
    for k in range(5):
        r[7 + 3 * k] = r[8 + 3 * k] = r[9 + 3 * k] = None
    for k, (a, b, c) in enumerate(ex):
        r[7 + 3 * k], r[8 + 3 * k], r[9 + 3 * k] = a, b, c


class Base:
    def __init__(self, path=FILE):
        self.path = path
        ws = openpyxl.load_workbook(path)['Vocab']
        self.header = [c.value for c in ws[1]]
        self.rows = [list(r)[:NCOL] + [None] * (NCOL - len(r)) for r in ws.iter_rows(min_row=2, values_only=True)]

    setex = staticmethod(setex)
    getex = staticmethod(getex)

    def fiche(self, entree):
        return [r for r in self.rows if r[E] == entree]

    def find(self, entree, mot, tr=None):
        """Ligne unique (Entrée, Mot), tr = début de la Traduction pour départager."""
        o = [r for r in self.rows if r[E] == entree and r[M] == mot and (tr is None or (r[T] or '').startswith(tr))]
        assert len(o) == 1, f'{len(o)} ligne(s) pour {entree!r} / {mot!r} / {tr!r}'
        return o[0]

    def row(self, entree, mot, nature, trad, syn=None, ant=None, fa=None, ex=(), zipf_of=None, assoc=False):
        """Nouvelle ligne ; zipf_of = mot distinctif sur lequel calculer la fréquence."""
        r = [None] * NCOL
        r[:7] = [entree, mot, nature, trad, syn, ant, fa]
        setex(r, list(ex))
        if zipf_of:
            r[ZIPF] = zipf(zipf_of)
            r[FREQ] = freq_class(r[ZIPF])
        r[TYPE] = 'association' if assoc else None
        return r

    def insert_after(self, ref, *new):
        i = self.rows.index(ref)
        self.rows[i + 1:i + 1] = list(new)

    def move(self, r, entree, after=None):
        """Déplace une ligne vers une autre fiche (à la fin, ou après la ligne `after`)."""
        self.rows.remove(r)
        r[E] = entree
        if after is None:
            last = max(i for i, x in enumerate(self.rows) if x[E] == entree)
            self.rows.insert(last + 1, r)
        else:
            self.insert_after(after, r)

    def add_entry(self, entree, rows):
        """Nouvelle fiche, insérée à sa place alphabétique."""
        assert not self.fiche(entree), 'la fiche existe déjà'
        names = []
        for r in self.rows:
            if r[E] not in names:
                names.append(r[E])
        nxt = next((n for n in sorted(names + [entree]) if n > entree), None)
        i = len(self.rows) if nxt is None else next(k for k, r in enumerate(self.rows) if r[E] == nxt)
        self.rows[i:i] = rows

    def check(self):
        issues, seen = [], set()
        for r in self.rows:
            for a, m, t in getex(r):
                if not m or m.lower() not in a.lower():
                    issues.append(f'mot à trouver absent : {r[M]} | {a} | {m}')
                if not t:
                    issues.append(f'traduction d\'exemple manquante : {r[M]} | {a}')
            if r[T] and ';' in r[T]:
                issues.append(f'« ; » dans Traduction : {r[M]}')
            k = (r[E], r[M], r[T])
            if k in seen:
                issues.append(f'doublon : {k}')
            seen.add(k)
            ex = [x[0] for x in getex(r)]
            if len(ex) != len(set(ex)):
                issues.append(f'exemple répété : {r[M]}')
        return issues

    def save(self, force=False):
        issues = self.check()
        if issues and not force:
            raise SystemExit('Contrôle en échec :\n' + '\n'.join(issues[:50]))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Vocab'
        ws.append(self.header)
        for c in ws[1]:
            c.font = Font(name='Calibri', size=11, bold=True)
        for r in self.rows:
            ws.append(r)
        wb.save(self.path)
        print(f'{len(self.rows)} lignes enregistrées, contrôle OK')


def show(rows):
    for r in rows:
        flags = ' [See also]' if r[TYPE] else ''
        print(f'{r[E]} | {r[M]} ({r[NAT]}) = {r[T]} | syn: {r[SYN] or "-"} | ant: {r[ANT] or "-"}'
              f' | FA: {r[FA] or "-"} | F{r[FREQ]} Z{r[ZIPF]}{flags}')
        for a, m, t in getex(r):
            print(f'    {a} [{m}] → {t}')


if __name__ == '__main__':
    cmd, args = (sys.argv[1], sys.argv[2:]) if len(sys.argv) > 1 else ('', [])
    B = Base()
    if cmd == 'show':
        for e in args:
            show(B.fiche(e))
    elif cmd == 'grep':
        rx = re.compile(args[0], re.I)
        show([r for r in B.rows if any(rx.search(str(r[c] or '')) for c in (M, T, SYN))])
    elif cmd == 'check':
        iss = B.check()
        print('\n'.join(iss) if iss else f'{len(B.rows)} lignes, contrôle OK')
    else:
        print(__doc__)
